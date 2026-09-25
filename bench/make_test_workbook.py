"""Write tests/sample_model.xlsx: a small model laid out differently from the reference workbook.

Labels in A, units in B, monthly timeline C:N with dates in row 3, an assumptions sheet, a summary
sheet, and a flat data table with no formulas.
"""
import os
from datetime import date

import openpyxl
from openpyxl.utils import get_column_letter as L

wb = openpyxl.Workbook()
a = wb.active
a.title = "Assumptions"
a.append(["Assumption", "Value", "Unit"])
a.append(["Monthly volume growth", 0.03, "%"])
a.append(["Unit price", 12.5, "$"])
a.append(["Unit cost", 7.0, "$"])
a.append(["Starting volume", 1000, "units"])

m = wb.create_sheet("Monthly")
m["A1"] = "Monthly forecast"
m["A3"], m["B3"] = "Month", "unit"
for i in range(12):
    m.cell(3, 3 + i, date(2025, 1 + i, 1))
labels = [("Volume", "units"), ("Revenue", "$"), ("Cost", "$"), ("Gross profit", "$")]
for r, (lab, unit) in enumerate(labels, start=5):
    m.cell(r, 1, lab)
    m.cell(r, 2, unit)
for i in range(12):
    c, p = L(3 + i), L(2 + i)
    m[f"{c}5"] = "=Assumptions!$B$5" if i == 0 else f"={p}5*(1+Assumptions!$B$2)"
    m[f"{c}6"] = f"={c}5*Assumptions!$B$3"
    m[f"{c}7"] = f"=-{c}5*Assumptions!$B$4"
    m[f"{c}8"] = f"={c}6+{c}7"

s = wb.create_sheet("Summary")
s.append(["Metric", "Value"])
s.append(["Total revenue", "=SUM(Monthly!C6:N6)"])
s.append(["Total gross profit", "=SUM(Monthly!C8:N8)"])
s.append(["Gross margin", "=B3/B2"])

d = wb.create_sheet("Sales data")
d.append(["Date", "Region", "Amount"])
for i in range(50):
    d.append([date(2025, 1 + i % 12, 1 + i % 28), ["North", "South", "East"][i % 3], 100 + i * 7])

os.makedirs("tests", exist_ok=True)
wb.save("tests/sample_model.xlsx")
print("wrote tests/sample_model.xlsx")
