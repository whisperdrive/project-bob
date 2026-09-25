"""Build a compact, row-oriented map of a formula-heavy workbook.

Usage: build_map.py <workbook.xlsx|.xlsm>

Outputs (in out/<workbook name>/):
  map.txt   - one line per line item: label, units, formula pattern(s) in R1C1, sample values
  model.db  - SQLite: sheets (detected layout), cells (every non-empty cell), rows (line items),
              edges (row -> row deps), names

Formulas come from openpyxl; cached values come from calamine (fast). Sheet layout (label / units
columns, timeline) is detected per sheet by layout.py.
"""
import json
import os
import re
import sqlite3
import sys
import time
from collections import defaultdict
from datetime import date, datetime

import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import column_index_from_string, get_column_letter
from python_calamine import CalamineWorkbook

from layout import describe, detect_layout

MAX_RANGE_ROWS = 300      # cap on row edges expanded from a single range ref
TABLE_MIN_ROWS = 50       # consecutive formula-free rows with the same columns are summarised as one table

CELL_RE = re.compile(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$")
COL_RE = re.compile(r"^(\$?)([A-Z]{1,3})$")
ROW_RE = re.compile(r"^(\$?)(\d+)$")


def split_sheet(ref: str) -> tuple[str | None, str]:
    if "!" in ref:
        sheet, addr = ref.rsplit("!", 1)
        return sheet.strip("'").replace("''", "'"), addr
    return None, ref


def cell_to_r1c1(addr: str, row: int, col: int) -> tuple[str, tuple[int, int] | None]:
    """Return R1C1 text for one endpoint and its absolute (row, col) target."""
    m = CELL_RE.match(addr)
    if m:
        cabs, cl, rabs, rn = m.groups()
        r, c = int(rn), column_index_from_string(cl)
        rtxt = f"R{r}" if rabs else ("R" if r == row else f"R[{r - row}]")
        ctxt = f"C{c}" if cabs else ("C" if c == col else f"C[{c - col}]")
        return rtxt + ctxt, (r, c)
    m = COL_RE.match(addr)
    if m:
        cabs, cl = m.groups()
        c = column_index_from_string(cl)
        return (f"C{c}" if cabs else ("C" if c == col else f"C[{c - col}]")), None
    m = ROW_RE.match(addr)
    if m:
        rabs, rn = m.groups()
        r = int(rn)
        return (f"R{r}" if rabs else ("R" if r == row else f"R[{r - row}]")), (r, None)
    return addr, None


def to_pattern(formula: str, row: int, col: int, sheet: str):
    """Convert an A1 formula to R1C1 text; also return referenced (sheet, row_lo, row_hi) spans and names."""
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return formula, [], []
    out, spans, names = ["="], [], []
    for t in tokens:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            tsheet, addr = split_sheet(t.value)
            parts = addr.split(":")
            conv = [cell_to_r1c1(p, row, col) for p in parts]
            if all(x[0] == p for x, p in zip(conv, parts)) and not CELL_RE.match(parts[0]):
                names.append(t.value)  # a defined name (or something we can't parse)
                out.append(t.value)
                continue
            prefix = t.value[: len(t.value) - len(addr)]
            out.append(prefix + ":".join(x[0] for x in conv))
            rows = [x[1][0] for x in conv if x[1]]
            if rows:
                spans.append((tsheet or sheet, min(rows), max(rows)))
        else:
            out.append(t.value)
    return "".join(out), spans, names


def fmt(v) -> str:
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e12:
            return str(int(v))
        return f"{v:.4g}"
    return str(v)[:40]


def out_dir(path: str) -> str:
    """out/<workbook file name without extension>, shared with tools.py and compare.py."""
    return os.path.join("out", os.path.splitext(os.path.basename(path))[0])


def summarize_table(run: list[tuple[int, str, dict]]) -> str:
    """One line for a block of same-shaped constant rows: header, row count, per-column type and range."""
    first, last = run[0][0], run[-1][0]
    header = None
    if all(isinstance(v, str) for v in run[0][2].values()) and \
            not all(isinstance(v, str) for v in run[1][2].values()):
        header, run = run[0][2], run[1:]
    cols = []
    for col in run[0][2]:
        vals = [rv[col] for _, _, rv in run if rv.get(col) not in ("", None)]
        name = f"{col}" + (f"={header[col]}" if header else "")
        if vals and all(isinstance(v, (date, datetime)) for v in vals):
            cols.append(f"{name} (date {min(vals)}..{max(vals)})")
        elif vals and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
            cols.append(f"{name} (number {fmt(min(vals))}..{fmt(max(vals))})")
        else:
            distinct = list(dict.fromkeys(fmt(v) for v in vals))
            eg = ", ".join(distinct[:5]) + (", ..." if len(distinct) > 5 else "")
            cols.append(f"{name} (text, {len(distinct)} distinct: {eg})")
    return (f"r{first}-r{last} TABLE {len(run)} data rows" + (" + header" if header else "")
            + " | " + "; ".join(cols))


def error_cells(path: str) -> dict[str, dict[tuple[int, int], str]]:
    """Cells whose saved result is an Excel error (#N/A, #REF!, #DIV/0!...), per sheet.

    calamine returns these as empty strings, so read them from the sheet XML (cells with t="e")."""
    import zipfile
    from xml.etree import ElementTree as ET
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    z = zipfile.ZipFile(path)
    rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
    out: dict[str, dict[tuple[int, int], str]] = {}
    cell_re = re.compile(rb'<c r="([A-Z]{1,3})(\d+)"[^>]*?t="e"[^>]*>.*?<v>([^<]*)</v>', re.S)
    for sh in ET.fromstring(z.read("xl/workbook.xml")).find("m:sheets", ns):
        target = rels.get(sh.get(f"{{{ns['r']}}}id"), "")
        target = target.lstrip("/") if target.startswith("/") else "xl/" + target
        if target not in z.namelist():
            continue
        errs = {(int(r), column_index_from_string(c.decode())): v.decode()
                for c, r, v in cell_re.findall(z.read(target))}
        if errs:
            out[sh.get("name")] = errs
    return out


def main(path: str, out: str | None = None, progress=None) -> dict:
    """Build map.txt + model.db in `out` (default out/<file stem>).

    progress(fraction 0..1, message) is called as sheets are read, for UIs.
    """
    if not path.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("build_map needs .xlsx or .xlsm (openpyxl can't read formulas from .xlsb/.xls); "
                         "save a copy from Excel first")
    report = progress or (lambda frac, msg: None)
    t0 = time.time()
    report(0.0, "Opening workbook")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cal = CalamineWorkbook.from_path(path)
    errors = error_cells(path)

    # defined names -> (sheet, row_lo, row_hi)
    name_targets: dict[str, list[tuple[str, int, int]]] = {}
    name_rows = []
    for n, dn in wb.defined_names.items():
        txt = dn.attr_text
        name_rows.append((n, txt))
        tsheet, addr = split_sheet(txt)
        if tsheet:
            rows = [cell_to_r1c1(p, 1, 1)[1] for p in addr.split(":")]
            rows = [r[0] for r in rows if r]
            if rows:
                name_targets[n.lower()] = [(tsheet, min(rows), max(rows))]

    out = out or out_dir(path)
    os.makedirs(out, exist_ok=True)
    db_path = os.path.join(out, "model.db")
    if os.path.exists(db_path):
        os.remove(db_path)
    db = sqlite3.connect(db_path)
    db.executescript("""
        CREATE TABLE sheets(sheet TEXT, state TEXT, layout TEXT, summary TEXT);
        CREATE TABLE cells(sheet TEXT, row INT, col INT, addr TEXT, formula TEXT, value);
        CREATE TABLE rows(sheet TEXT, row INT, section TEXT, label TEXT, units TEXT,
                          n_formula INT, n_const INT, patterns TEXT, samples TEXT);
        CREATE TABLE edges(src_sheet TEXT, src_row INT, dst_sheet TEXT, dst_row INT);
        CREATE TABLE names(name TEXT, ref TEXT);
    """)
    db.executemany("INSERT INTO names VALUES (?,?)", name_rows)

    lines = [f"WORKBOOK {os.path.basename(path)}",
             "Legend: sheet!row label [units] f=formula cells c=constant cells | pattern(s) in R1C1 "
             "(R[n]=relative row, R5=absolute row, C likewise) x count cols | eg sample cached values"]
    sheet_rows: dict[str, set[int]] = defaultdict(set)
    all_edges = set()

    n_sheets = len(wb.worksheets)
    for i_sheet, ws in enumerate(wb.worksheets):
        name = ws.title
        report(0.02 + 0.93 * i_sheet / n_sheets, f"Reading sheet {i_sheet + 1} of {n_sheets}: {name}")
        max_row, last_report = ws.max_row or 1, 0
        values = cal.get_sheet_by_name(name).to_python(skip_empty_area=False)
        sheet_errors = errors.get(name, {})
        val = lambda r, c: sheet_errors.get((r, c)) or (
            values[r - 1][c - 1] if r - 1 < len(values) and c - 1 < len(values[r - 1]) else None)
        lay = detect_layout(values)
        label_col, units_col = lay["label_col"], lay["units_col"]
        db.execute("INSERT INTO sheets VALUES (?,?,?,?)",
                   (name, ws.sheet_state, json.dumps(lay), describe(lay)))
        lines.append(f"\n## {name} ({ws.sheet_state}) - {describe(lay)}")
        section = ""
        cell_batch = []
        run: list[tuple[int, str, dict]] = []  # pending formula-free rows that may form a table

        def flush() -> None:
            if len(run) >= TABLE_MIN_ROWS:
                lines.append(summarize_table(run))
            else:
                lines.extend(line for _, line, _ in run)
            run.clear()

        for row in ws.iter_rows():
            cells = [c for c in row if c.value is not None]
            if not cells:
                continue
            r = cells[0].row
            if r - last_report >= 2000:
                last_report = r
                report(0.02 + 0.93 * (i_sheet + min(r / max_row, 1)) / n_sheets,
                       f"Reading sheet {i_sheet + 1} of {n_sheets}: {name} (row {r:,} of {max_row:,})")
            label_bits, units = [], ""
            patterns: dict[str, list[str]] = defaultdict(list)
            n_formula = n_const = 0
            samples = []
            first_text = None  # fallback label when the label column is empty on this row
            rowvals = {}
            for c in cells:
                v, cv = c.value, val(c.row, c.column)
                rowvals[c.column_letter] = cv
                is_f = isinstance(v, str) and v.startswith("=")
                cell_batch.append((name, r, c.column, c.coordinate, v if is_f else None,
                                   None if cv == "" else (cv if isinstance(cv, (int, float, str)) else str(cv))))
                # label area = label column plus anything left of it (e.g. section numbers); numbers in the
                # label or units column itself are data (summary blocks often put totals there)
                numeric = isinstance(cv, (int, float)) and not isinstance(cv, bool)
                in_label = label_col and (c.column < label_col or (c.column == label_col and not numeric))
                in_units = c.column == units_col and not numeric
                if in_label or in_units:
                    shown = cv if cv not in ("", None) else v
                    if in_units:
                        units = fmt(shown)
                    elif shown not in ("", None):
                        label_bits.append(fmt(shown))
                    continue
                if is_f:
                    n_formula += 1
                    pat, spans, names = to_pattern(v, r, c.column, name)
                    patterns[pat].append(c.coordinate)
                    for s, lo, hi in spans:
                        for tr in range(lo, min(hi, lo + MAX_RANGE_ROWS) + 1):
                            all_edges.add((name, r, s, tr))
                    for nm in names:
                        for s, lo, hi in name_targets.get(nm.split("!")[-1].lower(), []):
                            for tr in range(lo, min(hi, lo + MAX_RANGE_ROWS) + 1):
                                all_edges.add((name, r, s, tr))
                else:
                    n_const += 1
                if first_text is None and isinstance(cv, str) and cv.strip():
                    first_text = cv.strip()[:60]
                if cv not in ("", None, 0, 0.0) and len(samples) < 3:
                    samples.append(f"{c.column_letter}={fmt(cv)}")
            if len(cell_batch) > 50_000:
                db.executemany("INSERT INTO cells VALUES (?,?,?,?,?,?)", cell_batch)
                cell_batch.clear()
            label = " ".join(label_bits) or (first_text or "")
            if not (n_formula or n_const):
                if label:
                    section = label
                    flush()
                    lines.append(f"# r{r} {label}")
                continue
            sheet_rows[name].add(r)
            pats = sorted(patterns.items(), key=lambda kv: -len(kv[1]))
            pat_txt = "; ".join(
                f"{p} x{len(cs)}" + (f" ({cs[0]}..{cs[-1]})" if len(cs) > 1 else f" ({cs[0]})")
                for p, cs in pats[:3])
            if len(pats) > 3:
                pat_txt += f"; +{len(pats) - 3} more patterns"
            line = f"r{r} {label or '?'} [{units}] f={n_formula} c={n_const}"
            if pat_txt:
                line += f" | {pat_txt}"
            if samples:
                line += f" | eg {', '.join(samples)}"
            if n_formula == 0:
                if run and (r != run[-1][0] + 1 or rowvals.keys() != run[-1][2].keys()):
                    flush()
                run.append((r, line, rowvals))
            else:
                flush()
                lines.append(line)
            db.execute("INSERT INTO rows VALUES (?,?,?,?,?,?,?,?,?)",
                       (name, r, section, label, units, n_formula, n_const, pat_txt, ", ".join(samples)))
        flush()
        db.executemany("INSERT INTO cells VALUES (?,?,?,?,?,?)", cell_batch)

    # keep only edges that land on a real line item, and drop self-loops
    report(0.96, "Indexing dependencies")
    edges = [(a, b, c, d) for a, b, c, d in all_edges
             if d in sheet_rows.get(c, ()) and (a, b) != (c, d)]
    db.executemany("INSERT INTO edges VALUES (?,?,?,?)", edges)
    db.executescript("""
        CREATE INDEX ix_cells ON cells(sheet,row,col);
        CREATE INDEX ix_rows ON rows(sheet,row);
        CREATE INDEX ix_e1 ON edges(src_sheet,src_row);
        CREATE INDEX ix_e2 ON edges(dst_sheet,dst_row);
    """)
    db.commit()

    with open(os.path.join(out, "map.txt"), "w") as f:
        f.write("\n".join(lines) + "\n")
    n_rows = sum(len(v) for v in sheet_rows.values())
    print(f"{out}: line items={n_rows} edges={len(edges)} names={len(name_rows)} secs={time.time() - t0:.1f}")
    report(1.0, "Workbook database built")
    return {"out": out, "db": db_path, "sheets": n_sheets, "line_items": n_rows, "edges": len(edges),
            "secs": round(time.time() - t0, 1)}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    try:
        main(sys.argv[1])
    except ValueError as e:
        sys.exit(str(e))
